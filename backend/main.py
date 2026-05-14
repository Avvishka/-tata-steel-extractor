"""
Tata Steel PDF Extractor — FastAPI Backend
==========================================
Wraps the extraction logic as a REST API.
Deploy free on Render.com.
"""

import re
import cv2
import json
import base64
import time
import io
import os
import numpy as np
import pytesseract
import tempfile

from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from PIL import Image
from pdf2image import convert_from_bytes
from openai import OpenAI
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = FastAPI(title="Tata Steel PDF Extractor", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

GITHUB_TOKEN = os.environ.get("GITHUB_TOKEN", "")

PARAM_COLS = [
    "Page No", "Customer", "Test Standard", "Pipe/Coil Size",
    "Specification", "Test", "Specimen Type", "TPI Agency", "Tester",
]

RESULT_COLS = [
    "Pipe/Coil No", "Heat No", "Thk mm", "Width mm", "Area mm2",
    "MGL mm", "YL R0.5 MPa", "YS R0.5 MPa", "UTL kN",
    "UTS MPa", "A50mm %", "YS/UTS",
]

ALL_COLS = PARAM_COLS + RESULT_COLS

RESULTS_PROMPT = """
You are reading a scanned Tata Steel tensile test certificate image.

Find the Results table. It always has exactly 3 rows:
  ROW 1 = column headers (Pipe/Coil No, Heat No, Thk, Width, Area, ...)
  ROW 2 = units (mm, mm, mm2, kN, MPa ...) — IGNORE THIS ROW COMPLETELY
  ROW 3 = actual data values — EXTRACT ONLY THIS ROW

Two table types:
  FULL (TBT specimen): 12 columns including MGL, YL R0.5, YS R0.5, A50mm, YS/UTS
  SHORT (WTT specimen): 7 columns only: Pipe/Coil No, Heat No, Thk, Width, Area, UTL, UTS

STRICT RULES:
1. Decimal points exact: 6.65 not 665, 33.30 not 3330, 140.33 not 14033
2. Pipe/Coil No = N + 8 digits e.g. N01059417
3. Heat No = alphanumeric e.g. 26C14031
4. Numbers only — no units like mm MPa kN
5. Use "" for missing columns
6. Return ONLY valid JSON — no explanation no markdown

JSON to return:
{
  "Pipe/Coil No": "",
  "Heat No": "",
  "Thk mm": "",
  "Width mm": "",
  "Area mm2": "",
  "MGL mm": "",
  "YL R0.5 MPa": "",
  "YS R0.5 MPa": "",
  "UTL kN": "",
  "UTS MPa": "",
  "A50mm %": "",
  "YS/UTS": ""
}
"""


def preprocess_image(image):
    if len(image.shape) == 3:
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    else:
        gray = image.copy()
    gray = cv2.resize(gray, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC)
    gray = cv2.fastNlMeansDenoising(gray)
    kernel = np.array([[0, -1, 0], [-1, 5, -1], [0, -1, 0]])
    gray = cv2.filter2D(gray, -1, kernel)
    bw = cv2.adaptiveThreshold(
        gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 10
    )
    return bw


def extract_parameters(page_pil, page_no):
    image = np.array(page_pil)
    image = cv2.cvtColor(image, cv2.COLOR_RGB2BGR)
    proc = preprocess_image(image)
    text = pytesseract.image_to_string(proc, config="--oem 3 --psm 6")
    text_upper = text.upper()

    data = {col: "" for col in PARAM_COLS}
    data["Page No"] = page_no

    m = re.search(r"M/S\s*BGL", text, re.IGNORECASE)
    if m:
        data["Customer"] = "M/S BGL"

    m = re.search(r"ASTM\s*A\d+[:\-]?\d+[A-Z]?", text, re.IGNORECASE)
    if m:
        data["Test Standard"] = m.group(0).strip()

    m = re.search(
        r"\d+\.?\d*\s*MM\s*OD\s*X\s*\d+\.?\d*\s*MM\s*(?:THK|WT)", text, re.IGNORECASE
    )
    if m:
        data["Pipe/Coil Size"] = m.group(0).strip()

    m = re.search(r"API\s*5L.*?PSL[- ]?\d", text, re.IGNORECASE)
    if m:
        data["Specification"] = m.group(0).strip()

    if "TENSILE" in text_upper:
        data["Test"] = "TENSILE TEST"

    m = re.search(r"\b(TBT|TWT|WTT)\b", text_upper)
    if m:
        data["Specimen Type"] = m.group(1)

    if any(x in text_upper for x in ["EDLIPSE", "EDUPSE", "ECLIPSE"]):
        data["TPI Agency"] = "EDLIPSE"

    if "ABHISHEK" in text_upper:
        data["Tester"] = "ABHISHEK TANDEL"

    return data


def to_base64(pil_image):
    buf = io.BytesIO()
    pil_image.convert("RGB").save(buf, format="JPEG", quality=95)
    return base64.standard_b64encode(buf.getvalue()).decode("utf-8")


def extract_results(client, pil_image):
    empty = {col: "" for col in RESULT_COLS}
    for attempt in range(3):
        try:
            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:image/jpeg;base64,{to_base64(pil_image)}",
                                    "detail": "high",
                                },
                            },
                            {"type": "text", "text": RESULTS_PROMPT},
                        ],
                    }
                ],
                max_tokens=512,
                temperature=0,
            )
            raw = response.choices[0].message.content.strip()
            raw = re.sub(r"^```(?:json)?\s*", "", raw)
            raw = re.sub(r"\s*```$", "", raw).strip()
            data = json.loads(raw)

            result = {}
            for col in RESULT_COLS:
                val = str(data.get(col, "")).strip()
                val = re.sub(r"[|\[\]\\_]", "", val).strip()
                if col == "Pipe/Coil No":
                    val = re.sub(r"\s+", "", val)
                    val = re.sub(r"^N[OQ]", "N0", val, flags=re.IGNORECASE).upper()
                elif col == "Heat No":
                    val = re.sub(r"[+!]", "1", val)
                    val = re.sub(r"[€£]", "C", val)
                    val = re.sub(r"[^A-Za-z0-9]", "", val).upper()
                elif val:
                    nums = re.findall(r"\d+\.\d+|\d+", val)
                    val = nums[0] if nums else ""
                result[col] = val

            return result

        except json.JSONDecodeError:
            return empty
        except Exception as e:
            err = str(e)
            if "429" in err or "rate" in err.lower():
                time.sleep(30)
            else:
                return empty

    return empty


def build_excel(rows):
    df_combined = pd.DataFrame(rows, columns=ALL_COLS).fillna("")
    df_params = pd.DataFrame(rows, columns=PARAM_COLS).fillna("")
    df_results = pd.DataFrame(rows, columns=["Page No"] + RESULT_COLS).fillna("")

    output = io.BytesIO()
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_combined.to_excel(writer, index=False, sheet_name="Combined")
        df_params.to_excel(writer, index=False, sheet_name="Parameters")
        df_results.to_excel(writer, index=False, sheet_name="Results Table")

        for sheet_name in ["Combined", "Parameters", "Results Table"]:
            ws = writer.sheets[sheet_name]
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E79")
                cell.font = Font(color="FFFFFF", bold=True, size=10)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = border
            ws.row_dimensions[1].height = 35

            for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
                fill = (
                    PatternFill("solid", fgColor="EBF3FB")
                    if i % 2 == 0
                    else PatternFill()
                )
                for cell in row:
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center"
                    )
                    cell.border = border
                    cell.fill = fill

            for col_cells in ws.columns:
                w = max(
                    (len(str(c.value)) for c in col_cells if c.value), default=10
                )
                ws.column_dimensions[col_cells[0].column_letter].width = min(
                    w + 4, 25
                )
            ws.freeze_panes = "A2"

    output.seek(0)
    return output


@app.get("/")
def root():
    return {"status": "ok", "message": "Tata Steel PDF Extractor API"}


@app.get("/health")
def health():
    return {"status": "healthy"}


@app.post("/extract")
async def extract_pdf(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Only PDF files are accepted")

    token = GITHUB_TOKEN
    if not token:
        raise HTTPException(
            status_code=500,
            detail="GITHUB_TOKEN not configured on server"
        )

    client = OpenAI(
        base_url="https://models.inference.ai.azure.com",
        api_key=token,
    )

    pdf_bytes = await file.read()

    try:
        pages = convert_from_bytes(pdf_bytes, dpi=200)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF conversion failed: {str(e)}")

    all_rows = []

    for page_no, pil in enumerate(pages, start=1):
        params = extract_parameters(pil, page_no)
        results = extract_results(client, pil)
        row = {**params, **results}
        all_rows.append(row)
        if page_no < len(pages):
            time.sleep(2)

    return JSONResponse(content={"pages": len(pages), "rows": all_rows})


@app.post("/extract-excel")
async def extract_pdf_excel(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Only PDF files are accepted")

    token = GITHUB_TOKEN
    if not token:
        raise HTTPException(
            status_code=500,
            detail="GITHUB_TOKEN not configured on server"
        )

    client = OpenAI(
        base_url="https://models.inference.ai.azure.com",
        api_key=token,
    )

    pdf_bytes = await file.read()

    try:
        pages = convert_from_bytes(pdf_bytes, dpi=200)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF conversion failed: {str(e)}")

    all_rows = []
    for page_no, pil in enumerate(pages, start=1):
        params = extract_parameters(pil, page_no)
        results = extract_results(client, pil)
        row = {**params, **results}
        all_rows.append(row)
        if page_no < len(pages):
            time.sleep(2)

    excel_buffer = build_excel(all_rows)

    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=Tata_Steel_Extracted.xlsx"
        },
    )

